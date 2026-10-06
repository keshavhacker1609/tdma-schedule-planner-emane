"""Builds docs/TDMA_Planner_Presentation.pptx.

    CANDIDATE="Your Name" python docs/tools/build_deck.py
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "..", "figures")
OUT = os.path.join(HERE, "..", "TDMA_Planner_Presentation.pptx")
CANDIDATE = os.environ.get("CANDIDATE", "Candidate")

NAVY = RGBColor(0x14, 0x2b, 0x4d)
BLUE = RGBColor(0x1f, 0x6f, 0xeb)
GREEN = RGBColor(0x2f, 0x9e, 0x44)
RED = RGBColor(0xe0, 0x31, 0x31)
GREY = RGBColor(0x49, 0x50, 0x57)
LIGHT = RGBColor(0xf1, 0xf3, 0xf5)
WHITE = RGBColor(0xff, 0xff, 0xff)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, content, size=18, bold=False, color=GREY, align=PP_ALIGN.LEFT,
         font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
        p.space_after = Pt(6)
    return box


def bullets(slide, x, y, w, items, size=19):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.5 * len(items) + 1))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = "•  " + item
        r.font.size = Pt(size)
        r.font.color.rgb = GREY
        r.font.name = "Calibri"
        p.space_after = Pt(10)
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def label(slide, x, y, w, h, content, fill, color=WHITE, size=16, bold=True):
    s = rect(slide, x, y, w, h, fill)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = content if isinstance(content, list) else [content]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size if i == 0 else size - 3)
        r.font.bold = bold if i == 0 else False
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return s


def slide(title, notes, n):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 0.12, BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.6, 0.35, 12, 0.9, title, size=32, bold=True, color=NAVY)
    text(s, 12.2, 7.0, 0.8, 0.3, str(n), size=11, color=GREY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def pic(s, name, x, y, w):
    return s.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), width=Inches(w))


# 1 ------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 2.0, 11.5, 1.2, "TDMA Schedule Planner and Optimizer", size=44, bold=True, color=WHITE)
text(s, 0.9, 3.2, 11.5, 0.8, "Distance-2 graph colouring with EMANE integration", size=24,
     color=RGBColor(0xc5, 0xd5, 0xf0))
text(s, 0.9, 5.2, 11.5, 0.5, f"{CANDIDATE}  |  Wireless Protocol Development Internship task  |  Vaan Megam Networks",
     size=16, color=RGBColor(0xc5, 0xd5, 0xf0))
s.notes_slide.notes_text_frame.text = (
    "Introduce the task in one sentence: plan collision-free time slots for radios sharing one frequency, "
    "reuse slots where radios are far apart, and enforce the schedule in EMANE.")

# 2 ------------------------------------------------------------- problem
s = slide("The problem", "Explain TDMA: one frequency, radios take turns. Close radios must not share a slot; far radios can. "
          "The goal is the shortest repeating frame with no collisions. Mention that the minimum is NP-hard so heuristics are needed.", 2)
bullets(s, 0.7, 1.5, 7.3, [
    "Radios share one frequency and transmit in numbered time slots",
    "Radios close together must use different slots, or they interfere",
    "Radios far apart may reuse a slot: spatial reuse",
    "Goal: fewest slots per frame, zero collisions",
    "Minimum slot count is NP-hard, so heuristics plus a quality bound",
], size=20)
label(s, 8.6, 1.7, 4.2, 1.0, ["Part 1 (mandatory)", "Python scheduler: graph model + colouring"], BLUE, size=19)
label(s, 8.6, 3.1, 4.2, 1.0, ["Part 2 (bonus)", "EMANE TDMA emulation of the schedule"], NAVY, size=19)
label(s, 8.6, 4.5, 4.2, 1.0, ["Documentation", "Design, thought process, final values"], GREY, size=19)

# 3 ------------------------------------------------------------- hidden terminal
s = slide("Why two hops, not one", "A and C cannot hear each other but both reach B, so their frames collide at B. "
          "A one-hop rule misses this. This is the hidden-terminal problem and the reason for distance-2 colouring.", 3)
pic(s, "hidden_terminal.png", 0.7, 1.5, 7.2)
bullets(s, 0.7, 3.9, 7.4, [
    "Direct link (1 hop): A and B must differ",
    "Hidden terminal (2 hops): A and C must differ",
    "3 or more hops apart: free to share a slot",
], size=20)
label(s, 8.7, 1.6, 4.1, 1.4, ["Conflict graph = G squared", "edge if 1 or 2 hops apart in G"], BLUE, size=20)
label(s, 8.7, 3.3, 4.1, 1.4, ["Slot = colour", "proper colouring of G squared"], NAVY, size=20)
label(s, 8.7, 5.0, 4.1, 1.4, ["Spatial reuse", "comes from the model itself"], GREEN, size=20)

# 4 ------------------------------------------------------------- pipeline
s = slide("Optimisation pipeline", "Each stage keeps the best colouring so far. The clique lower bound tells us when we are "
          "provably optimal. Be ready to explain DSATUR (colour the vertex that sees most distinct neighbour colours), "
          "TabuCol (fix k = best-1 colours, move conflicting nodes, tabu list), and the exact search.", 4)
steps = [("Clique bound", "lower bound"), ("DSATUR + 2 greedy", "constructive"), ("Random restarts", "300 x DSATUR"),
         ("TabuCol", "remove one colour"), ("Exact search", "prove or improve")]
for i, (a, b) in enumerate(steps):
    label(s, 0.5 + i * 2.55, 2.0, 2.3, 1.3, [a, b], BLUE if i < 4 else NAVY, size=17)
    if i < 4:
        text(s, 2.78 + i * 2.55, 2.25, 0.4, 0.6, "▶", size=22, color=GREY, align=PP_ALIGN.CENTER)
bullets(s, 0.7, 3.9, 12, [
    "Stops as soon as colours = lower bound: result is provably optimal",
    "Otherwise the report prints the interval [lower bound, result]: never over-claims",
    "Seeded randomness: every run is reproducible",
], size=20)

# 5 ------------------------------------------------------------- verification
s = slide("Correctness: an independent checker", "The verifier shares no code with the optimizer, so a bug cannot hide itself. "
          "Mention 29 tests: boundary at exactly 500 m, isolated nodes, full mesh, 40 random topologies, bad input, CLI.", 5)
bullets(s, 0.7, 1.5, 7.2, [
    "verify.py works only from coordinates",
    "For every pair in the same slot: in range? share a neighbour in range?",
    "Reports the exact offending pair and reason, exit code 1",
    "\"Verified conflict-free\" printed only after the check passes",
    "29 automated tests including 40 random topologies",
], size=20)
label(s, 8.6, 1.7, 4.2, 1.2, ["29 tests", "all passing"], GREEN, size=24)
label(s, 8.6, 3.2, 4.2, 1.2, ["0 collisions", "in every verified schedule"], GREEN, size=24)
label(s, 8.6, 4.7, 4.2, 1.2, ["Strict input checks", "bad JSON, NaN, wrong arity"], NAVY, size=22)

# 6 ------------------------------------------------------------- results part 1
s = slide("Part 1 results (16 nodes, 500 m)", "The grid reaches the lower bound of 9, so it cannot be improved. Slot 0 carries four "
          "radios at once. Point at colours in the figure: same colour means same slot.", 6)
pic(s, "grid_slots.png", 0.6, 1.3, 5.6)
rows = [["Topology", "Links", "Slots", "Bound", "Status"],
        ["4x4 grid, 300 m", "42", "9", "9", "optimal"],
        ["Random layout", "39", "10", "10", "optimal"],
        ["Two clusters + isolated", "47", "8", "8", "optimal"]]
tbl = s.shapes.add_table(4, 5, Inches(6.7), Inches(1.7), Inches(6.2), Inches(2.2)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(15)
        p.font.bold = r == 0
        p.font.name = "Calibri"
text(s, 6.7, 4.3, 6.2, 2.0, ["Slot-level replay of the grid (84 links):", "distance-2 schedule: 84 delivered, 0 collided",
                             "plain distance-1 colouring: 20 delivered, 64 collided"], size=17)

# 7 ------------------------------------------------------------- architecture
s = slide("Part 2: from schedule to EMANE", "The bridge is the only link between the parts and contains no scheduling logic; "
          "it also refuses to export a schedule that breaks the two-hop rule. One container, one namespace per radio.", 7)
boxes = [("Coordinates", "JSON"), ("Python optimizer", "Part 1"), ("schedule.json", "contract"),
         ("bridge.py", "XML + EEL"), ("EMANE 1.5.3", "16 radios")]
for i, (a, b) in enumerate(boxes):
    label(s, 0.5 + i * 2.55, 1.9, 2.3, 1.3, [a, b], BLUE if i in (1, 4) else NAVY, size=17)
    if i < 4:
        text(s, 2.78 + i * 2.55, 2.15, 0.4, 0.6, "▶", size=22, color=GREY, align=PP_ALIGN.CENTER)
bullets(s, 0.7, 3.8, 12, [
    "One network namespace and one emane process per radio, in a single Docker container",
    "emaneevent-tdmaschedule delivers each NEM its own slots",
    "Pathloss events (80 dB in range, 250 dB beyond) reproduce the 500 m disk model exactly",
    "Default slot length 1 ms as specified; frame = number of colours",
], size=19)

# 8 ------------------------------------------------------------- schedule xml
s = slide("Schedule mapping", "Each node transmits in its colour's slot and listens in every other slot. Nodes with the same colour "
          "appear in one nodes list: that is spatial reuse expressed in EMANE's format.", 8)
code = rect(s, 0.7, 1.5, 6.9, 4.4, LIGHT)
box = text(s, 0.9, 1.6, 6.6, 4.2, [
    '<structure frames="1" slots="9"',
    '    slotduration="1000" bandwidth="1M"/>',
    '<multiframe frequency="2.4G" ...>',
    '  <frame index="0">',
    '    <slot index="0" nodes="1,4,13,16">',
    '      <tx/>',
    '    </slot>',
    '    <slot index="0" nodes="2,3,5,6,...">',
    '      <rx/>',
    '    </slot>',
    '    ...',
], size=15, color=NAVY, font="Consolas")
bullets(s, 8.0, 1.6, 4.9, [
    "Slot 0: Node_01, 04, 13, 16 transmit together",
    "Everyone else listens in slot 0",
    "1 ms slots, one frame, one frequency",
    "Test checks XML against the Python schedule",
], size=18)

# 9 ------------------------------------------------------------- emane results
s = slide("EMANE results", "Real run on EMANE 1.5.3. 120 of 120 pairs behaved per the 500 m topology at 10 ms slots. Under 16 simultaneous "
          "flows the distance-2 schedule lost 0.8 percent against 43.1 percent for the naive one. At 1 ms slots, laptop containers "
          "without realtime scheduling miss slot boundaries: 92 of 120. Say this openly.", 9)
pic(s, "emane_loss.png", 0.6, 1.4, 6.2)
bullets(s, 7.1, 1.5, 5.9, [
    "120 / 120 node pairs follow the 500 m topology",
    "NEM 1 neighbour table: exactly NEMs 2, 5, 6",
    "This schedule: 0.8 % loss (9 slots)",
    "Plain distance-1: 43.1 % loss (4 slots)",
    "1 ms slots need realtime scheduling: 92 / 120 on a laptop, so tests ran at 10 ms",
], size=18)

# 10 ------------------------------------------------------------ lessons
s = slide("What the real run taught me", "Be ready to discuss each of these: writing configuration from documentation was not enough; EMANE "
          "rejected it several times and each failure was fixed.", 10)
bullets(s, 0.7, 1.5, 12, [
    "PHY profile needs a subid parameter",
    "EEL timestamps reject -Inf; use 0.0",
    "emane --pidfile (not -p, which is realtime priority)",
    "Event service resolves scenario.eel relative to its working directory",
    "Slot length must match what the host can time; the schedule logic is unchanged",
], size=22)

# 11 ------------------------------------------------------------ limits
s = slide("Limitations and next steps", "Show you understand the model's boundaries: unit-disk range, equal airtime, static positions, "
          "clique bound not always tight.", 11)
bullets(s, 0.7, 1.5, 12, [
    "Unit-disk range model; real RF has fading and capture",
    "One slot per node: weighted colouring for per-node demand",
    "Static positions: incremental recolouring for mobility",
    "Multi-frequency schedules to shorten the frame further",
    "Faithful 1 ms emulation on a realtime-capable host",
], size=22)

# 12 ------------------------------------------------------------ demo
s = slide("Live demo plan", "1) run the CLI on the grid, 2) show matrix and slot reuse, 3) show slot_sim comparison, 4) show bridge output, "
          "5) show EMANE log. Keep a recorded fallback in case Docker is slow.", 12)
bullets(s, 0.7, 1.5, 7.5, [
    "python schedule_optimizer.py --nodes-file examples/grid_16.json",
    "python emane/slot_sim.py examples/grid_16.schedule.json",
    "python emane/bridge.py ... then inspect tdmaschedule.xml",
    "SLOT_US=10000 sh emane/run_in_docker.sh examples/grid_16.json",
    "python -m pytest -q",
], size=17)
label(s, 8.7, 1.8, 4.1, 1.4, ["Summary", "Part 1 complete and verified"], BLUE, size=22)
label(s, 8.7, 3.5, 4.1, 1.4, ["Part 2 executed", "on EMANE 1.5.3"], GREEN, size=22)
label(s, 8.7, 5.2, 4.1, 1.4, ["Documentation", "day-wise report + README"], NAVY, size=22)

prs.save(OUT)
print("wrote", OUT)
